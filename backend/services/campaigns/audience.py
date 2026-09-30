"""Spreadsheet audience. Headers first, phone column chosen by the caller."""
import csv
import io
import zipfile

from backend.services.campaigns.constants import FIELD_CAP, ROW_CAP, ZIP_UNCOMPRESSED_CAP
from backend.services.silence.phones import canonical


def preview(filename: str, payload: bytes) -> tuple[list[str], list[dict]]:
    headers, rows = read_table(filename, payload)
    return headers, rows[:5]


def read_table(filename: str, payload: bytes) -> tuple[list[str], list[dict]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        _reject_zip_bomb(payload)
        return _xlsx(payload)
    if name.endswith(".csv"):
        return _csv(payload)
    raise ValueError("csv_or_xlsx")


def recipients(rows: list[dict], phone_column: str) -> tuple[list[dict], int]:
    if phone_column not in (rows[0] if rows else {}):
        raise ValueError("phone_column")
    kept: list[dict] = []
    seen: set[str] = set()
    invalid = 0
    for index, row in enumerate(rows):
        if len(kept) >= ROW_CAP:
            break
        number = canonical(str(row.get(phone_column) or ""))
        if not number:
            invalid += 1
            continue
        if number in seen:
            continue
        seen.add(number)
        fields = {
            key: _cell(value)
            for key, value in row.items()
            if key != phone_column
        }
        kept.append({"phone": number, "fields": fields, "sort_order": index})
    return kept, invalid


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()[:FIELD_CAP]


def _reject_zip_bomb(payload: bytes) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(payload))
    except zipfile.BadZipFile as error:
        raise ValueError("bad_xlsx") from error
    total = sum(info.file_size for info in archive.infolist())
    if total > ZIP_UNCOMPRESSED_CAP:
        raise ValueError("xlsx_too_large")


def _xlsx(payload: bytes) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("bad_xlsx") from error
    try:
        sheet = book.worksheets[0]
        grid = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    if not grid:
        raise ValueError("empty")
    headers = [_header(cell, index) for index, cell in enumerate(grid[0])]
    rows = [_mapped(headers, line) for line in grid[1:] if any(cell not in (None, "") for cell in line)]
    return headers, rows


def _csv(payload: bytes) -> tuple[list[str], list[dict]]:
    text = _decode(payload)
    reader = csv.reader(io.StringIO(text))
    try:
        header_line = next(reader)
    except StopIteration as error:
        raise ValueError("empty") from error
    headers = [_header(cell, index) for index, cell in enumerate(header_line)]
    rows = []
    for line in reader:
        if not any(cell.strip() for cell in line):
            continue
        rows.append(_mapped(headers, line))
    return headers, rows


def _mapped(headers: list[str], line) -> dict:
    values = list(line)
    return {headers[index]: _cell(values[index] if index < len(values) else "") for index in range(len(headers))}


def _header(cell, index: int) -> str:
    text = str(cell or "").strip()
    return text or f"עמודה {index + 1}"


def _decode(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1255", "latin-1"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("latin-1", errors="ignore")
