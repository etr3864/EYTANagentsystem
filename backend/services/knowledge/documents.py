"""Document processing — PDF, DOCX, TXT parsing, chunking, and source edits."""
import io
from sqlalchemy.orm import Session
from sqlalchemy import select, delete as sql_delete

from backend.models.knowledge import Document, DocumentChunk
from . import embeddings
from .retrieval import (
    CANDIDATE_MULTIPLIER,
    DEFAULT_SEARCH_LIMIT,
    MAX_CHUNKS_PER_DOCUMENT,
    MAX_COSINE_DISTANCE,
    cap_per_key,
    like_contains,
    match_by_name,
)
from backend.core.logger import log_upload

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
MAX_SOURCE_CHARS = 200_000
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_TITLE_LEN = 255
MAX_BULK_DELETE = 100
SOURCE_TOO_LONG = (
    "המסמך ארוך מדי (עד 200,000 תווים, בערך 100 עמודים). "
    "פצל אותו לכמה מסמכים קצרים יותר והעלה כל אחד בנפרד, או קצר את הטקסט."
)
FILE_TOO_HEAVY = (
    "הקובץ גדול מדי (עד 10MB). "
    "דחוס את ה-PDF, או פצל אותו לקבצים קטנים יותר והעלה כל חלק כמסמך נפרד."
)


def _extract_text_pdf(content: bytes) -> str:
    import fitz

    doc = fitz.open(stream=content, filetype="pdf")
    try:
        return "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def _extract_text_docx(content: bytes) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_text_txt(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1255"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("קובץ הטקסט חייב להיות UTF-8")


def _extract_text(filename: str, content: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return _extract_text_pdf(content)
    if ext in ("docx", "doc"):
        return _extract_text_docx(content)
    if ext == "txt":
        return _extract_text_txt(content)
    raise ValueError(f"סוג קובץ לא נתמך: {ext}")


def _chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = text.strip()
    if len(text) <= chunk_size:
        return [text] if text else []

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        if end < len(text):
            for sep in ["\n\n", "\n", ". ", "! ", "? "]:
                last_sep = chunk.rfind(sep)
                if last_sep > chunk_size // 2:
                    chunk = chunk[: last_sep + len(sep)]
                    end = start + len(chunk)
                    break
        chunks.append(chunk.strip())
        start = end - overlap
    return [c for c in chunks if c]


def require_title(title: str) -> str:
    title = (title or "").strip()
    if not title:
        raise ValueError("חובה לתת כותרת למסמך")
    if len(title) > MAX_TITLE_LEN:
        raise ValueError(f"כותרת ארוכה מדי (עד {MAX_TITLE_LEN} תווים)")
    return title


def require_source(text: str) -> str:
    text = (text or "").strip()
    if not text:
        raise ValueError("אין תוכן במסמך")
    if len(text) > MAX_SOURCE_CHARS:
        raise ValueError(SOURCE_TOO_LONG)
    return text


def to_list_item(doc: Document) -> dict:
    created = doc.created_at.isoformat() if doc.created_at else None
    return {
        "id": doc.id,
        "filename": doc.filename,
        "file_type": doc.file_type,
        "file_size": doc.file_size,
        "chunk_count": doc.chunk_count,
        "has_source": bool(doc.source_text),
        "created_at": created,
    }


def to_detail(doc: Document) -> dict:
    item = to_list_item(doc)
    item["source_text"] = doc.source_text
    return item


def get_for_agent(db: Session, agent_id: int, doc_id: int) -> Document | None:
    return db.scalar(
        select(Document).where(Document.id == doc_id, Document.agent_id == agent_id)
    )


def get_by_agent(db: Session, agent_id: int) -> list[Document]:
    return list(db.scalars(
        select(Document)
        .where(Document.agent_id == agent_id, Document.is_active == True)
        .order_by(Document.created_at.desc())
    ))


def _replace_chunks(db: Session, doc: Document, text: str) -> None:
    db.execute(sql_delete(DocumentChunk).where(DocumentChunk.document_id == doc.id))
    db.flush()
    chunks = _chunk_text(text)
    if not chunks:
        raise ValueError("No text content found in document")
    embs = embeddings.get_embeddings_batch(chunks)
    for i, (chunk_text, emb) in enumerate(zip(chunks, embs)):
        db.add(DocumentChunk(
            document_id=doc.id,
            content=chunk_text,
            chunk_index=i,
            embedding=emb,
        ))
    doc.chunk_count = len(chunks)
    doc.source_text = text


def _new_document(db: Session, agent_id: int, title: str, file_type: str, file_size: int, text: str) -> Document:
    doc = Document(
        agent_id=agent_id,
        filename=title,
        file_type=file_type,
        file_size=file_size,
        chunk_count=0,
    )
    db.add(doc)
    db.flush()
    _replace_chunks(db, doc, text)
    db.commit()
    db.refresh(doc)
    log_upload("document", title, f"{doc.chunk_count} chunks")
    return doc


def upload(db: Session, agent_id: int, filename: str, content: bytes, title: str) -> Document:
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError(FILE_TOO_HEAVY)
    title = require_title(title)
    ext = filename.rsplit(".", 1)[-1].lower()
    text = require_source(_extract_text(filename, content))
    return _new_document(db, agent_id, title, ext, len(content), text)


def create_from_text(db: Session, agent_id: int, title: str, content: str) -> Document:
    title = require_title(title)
    text = require_source(content)
    return _new_document(db, agent_id, title, "txt", len(text.encode("utf-8")), text)


def update(
    db: Session,
    agent_id: int,
    doc_id: int,
    title: str | None = None,
    content: str | None = None,
) -> Document | None:
    doc = get_for_agent(db, agent_id, doc_id)
    if not doc:
        return None
    if title is not None:
        doc.filename = require_title(title)
    if content is not None:
        if doc.source_text is None:
            raise ValueError("אין טקסט מקור למסמך הזה — העלה אותו מחדש כדי לערוך")
        text = require_source(content)
        if text != doc.source_text:
            _replace_chunks(db, doc, text)
            doc.file_size = len(text.encode("utf-8"))
    db.commit()
    db.refresh(doc)
    return doc


def delete(db: Session, agent_id: int, doc_id: int) -> bool:
    doc = get_for_agent(db, agent_id, doc_id)
    if not doc:
        return False
    db.delete(doc)
    db.commit()
    return True


def bulk_delete(db: Session, agent_id: int, ids: list[int]) -> int:
    unique = list(dict.fromkeys(i for i in ids if isinstance(i, int) and i > 0))
    unique = unique[:MAX_BULK_DELETE]
    if not unique:
        return 0
    docs = list(db.scalars(
        select(Document).where(Document.agent_id == agent_id, Document.id.in_(unique))
    ))
    for doc in docs:
        db.delete(doc)
    db.commit()
    return len(docs)


def _resolve_document_ids(
    db: Session, agent_id: int, document: str | None
) -> tuple[list[int] | None, int]:
    docs = get_by_agent(db, agent_id)
    if not document or not document.strip():
        return None, len(docs)
    matched_names = set(match_by_name([d.filename for d in docs], document))
    return [d.id for d in docs if d.filename in matched_names], len(docs)


def _search_lexical(
    db: Session,
    agent_id: int,
    query: str,
    doc_ids: list[int] | None,
    fetch: int,
) -> list[dict]:
    if len(query.strip()) < 2:
        return []
    conditions = [
        Document.agent_id == agent_id,
        Document.is_active == True,
        DocumentChunk.content.ilike(like_contains(query), escape="\\"),
    ]
    if doc_ids is not None:
        conditions.append(Document.id.in_(doc_ids))
    rows = db.execute(
        select(DocumentChunk, Document)
        .join(Document)
        .where(*conditions)
        .order_by(DocumentChunk.chunk_index)
        .limit(fetch)
    ).all()
    return [
        {
            "document": doc.filename,
            "content": chunk.content,
            "chunk_index": chunk.chunk_index,
        }
        for chunk, doc in rows
    ]


def _search_semantic(
    db: Session,
    agent_id: int,
    query: str,
    doc_ids: list[int] | None,
    fetch: int,
    apply_cutoff: bool,
) -> list[dict]:
    query_embedding = embeddings.get_embedding(query)
    distance = DocumentChunk.embedding.cosine_distance(query_embedding)
    conditions = [
        Document.agent_id == agent_id,
        Document.is_active == True,
        DocumentChunk.embedding.isnot(None),
    ]
    if doc_ids is not None:
        conditions.append(Document.id.in_(doc_ids))
    rows = db.execute(
        select(DocumentChunk, Document, distance.label("distance"))
        .join(Document)
        .where(*conditions)
        .order_by(distance)
        .limit(fetch)
    ).all()
    scored = []
    for chunk, doc, dist in rows:
        if dist is None:
            continue
        if apply_cutoff and dist > MAX_COSINE_DISTANCE:
            continue
        scored.append({
            "document": doc.filename,
            "content": chunk.content,
            "chunk_index": chunk.chunk_index,
        })
    return scored


def search(
    db: Session,
    agent_id: int,
    query: str,
    limit: int = DEFAULT_SEARCH_LIMIT,
    document: str | None = None,
) -> list[dict]:
    if not query.strip():
        return []

    doc_ids, total_docs = _resolve_document_ids(db, agent_id, document)
    if document and document.strip() and doc_ids is not None and not doc_ids:
        return []
    if doc_ids is None and len(query.strip()) >= 4:
        auto_ids, auto_total = _resolve_document_ids(db, agent_id, query)
        if auto_ids and len(auto_ids) <= 2:
            doc_ids, total_docs = auto_ids, auto_total

    scoped = total_docs <= 1 or (doc_ids is not None and len(doc_ids) == 1)
    fetch = max(limit * CANDIDATE_MULTIPLIER, 20)
    per_doc = limit if scoped else MAX_CHUNKS_PER_DOCUMENT

    found = _search_lexical(db, agent_id, query, doc_ids, fetch)
    if not found:
        found = _search_semantic(db, agent_id, query, doc_ids, fetch, apply_cutoff=not scoped)
    return cap_per_key(found, "document", per_doc, limit)
