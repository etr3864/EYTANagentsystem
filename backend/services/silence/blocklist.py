"""Manual numbers and customer opt-outs. One row per agent and phone."""
import io
import re
from datetime import datetime

from sqlalchemy.orm import Session

from backend.models.blocked_number import BlockedNumber
from backend.services.silence.phones import canonical

_CHUNK = re.compile(r"\+?\d[\d\s\-()]{6,}\d")
_MAX_NUMBERS = 20000
PAGE_SIZE = 50
_QUOTE_CAP = 240


def page(db: Session, agent_id: int, page_number: int) -> dict:
    page_number = max(page_number, 1)
    query = db.query(BlockedNumber).filter(BlockedNumber.agent_id == agent_id)
    total = query.count()
    rows = (
        query.order_by(BlockedNumber.phone)
        .offset((page_number - 1) * PAGE_SIZE)
        .limit(PAGE_SIZE)
        .all()
    )
    return {
        "items": [_row(row) for row in rows],
        "total": total,
        "page": page_number,
        "page_size": PAGE_SIZE,
    }


def add_one(db: Session, agent_id: int, raw: str) -> str:
    number = canonical(raw)
    if not number:
        raise ValueError("invalid_phone")
    _insert(db, agent_id, [number])
    return number


def replace_one(db: Session, agent_id: int, old: str, new: str) -> str:
    current = canonical(old)
    number = canonical(new)
    if not current or not number:
        raise ValueError("invalid_phone")
    row = (
        db.query(BlockedNumber)
        .filter(BlockedNumber.agent_id == agent_id, BlockedNumber.phone == current)
        .first()
    )
    if row is None:
        raise ValueError("missing_phone")
    if number != current:
        taken = (
            db.query(BlockedNumber.id)
            .filter(BlockedNumber.agent_id == agent_id, BlockedNumber.phone == number)
            .first()
        )
        if taken is not None:
            raise ValueError("duplicate_phone")
        row.phone = number
    return number


def remove(db: Session, agent_id: int, raw: str) -> None:
    clear_manual(db, agent_id, raw)


def remove_many(db: Session, agent_id: int, raw_phones: list[str]) -> None:
    numbers = _unique(raw_phones)
    if len(numbers) > _MAX_NUMBERS:
        raise ValueError("too_many")
    for number in numbers:
        clear_manual(db, agent_id, number)


def remove_all(db: Session, agent_id: int) -> None:
    db.query(BlockedNumber).filter(
        BlockedNumber.agent_id == agent_id,
        BlockedNumber.manual.is_(True),
        BlockedNumber.opted_out.is_(False),
    ).delete(synchronize_session=False)
    db.query(BlockedNumber).filter(
        BlockedNumber.agent_id == agent_id,
        BlockedNumber.manual.is_(True),
        BlockedNumber.opted_out.is_(True),
    ).update({BlockedNumber.manual: False}, synchronize_session=False)


def clear_manual(db: Session, agent_id: int, raw: str) -> None:
    row = _row_for(db, agent_id, raw)
    if row is None or not row.manual:
        return
    if row.opted_out:
        row.manual = False
        return
    db.delete(row)


def clear_opt_out(db: Session, agent_id: int, raw: str) -> None:
    row = _row_for(db, agent_id, raw)
    if row is None or not row.opted_out:
        return
    row.opted_out = False
    row.quote = None
    row.opted_out_at = None
    if not row.manual:
        db.delete(row)


def mark_opted_out(db: Session, agent_id: int, raw: str, quote: str) -> None:
    number = canonical(raw)
    if not number:
        return
    row = (
        db.query(BlockedNumber)
        .filter(BlockedNumber.agent_id == agent_id, BlockedNumber.phone == number)
        .first()
    )
    if row is None:
        row = BlockedNumber(agent_id=agent_id, phone=number, manual=False, opted_out=True)
        db.add(row)
    row.opted_out = True
    row.quote = (quote or "")[:_QUOTE_CAP] or None
    row.opted_out_at = datetime.utcnow()


def _delete_phones(db: Session, agent_id: int, numbers: list[str]) -> None:
    step = 500
    for start in range(0, len(numbers), step):
        db.query(BlockedNumber).filter(
            BlockedNumber.agent_id == agent_id,
            BlockedNumber.phone.in_(numbers[start:start + step]),
        ).delete(synchronize_session=False)


def import_numbers(db: Session, agent_id: int, raw_text: str, *, replace: bool) -> int:
    numbers = parse_numbers(raw_text)
    if len(numbers) > _MAX_NUMBERS:
        raise ValueError("too_many")
    if replace:
        remove_all(db, agent_id)
    _insert(db, agent_id, numbers)
    return len(numbers)


def parse_numbers(raw_text: str) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()
    for match in _CHUNK.findall(raw_text or ""):
        number = canonical(match)
        if number and number not in seen:
            seen.add(number)
            found.append(number)
    return found


def text_from_upload(filename: str, payload: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        raise ValueError("שמור את האקסל כ-xlsx או csv")
    if name.endswith(".xlsx"):
        return _xlsx_text(payload)
    for encoding in ("utf-8-sig", "cp1255", "latin-1"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


def _xlsx_text(payload: bytes) -> str:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("לא הצלחנו לקרוא את האקסל") from error
    parts: list[str] = []
    try:
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        parts.append(str(cell))
    finally:
        book.close()
    return "\n".join(parts)


def _insert(db: Session, agent_id: int, numbers: list[str]) -> None:
    if not numbers:
        return
    existing = {
        row.phone: row
        for row in db.query(BlockedNumber)
        .filter(BlockedNumber.agent_id == agent_id, BlockedNumber.phone.in_(numbers))
        .all()
    }
    for number in numbers:
        row = existing.get(number)
        if row is None:
            db.add(BlockedNumber(agent_id=agent_id, phone=number, manual=True))
            continue
        row.manual = True


def _row_for(db: Session, agent_id: int, raw: str) -> BlockedNumber | None:
    number = canonical(raw)
    if not number:
        return None
    return (
        db.query(BlockedNumber)
        .filter(BlockedNumber.agent_id == agent_id, BlockedNumber.phone == number)
        .first()
    )


def _unique(raw_phones: list[str]) -> list[str]:
    numbers: list[str] = []
    seen: set[str] = set()
    for raw in raw_phones:
        number = canonical(raw)
        if number and number not in seen:
            seen.add(number)
            numbers.append(number)
    return numbers


def _row(row: BlockedNumber) -> dict:
    return {
        "phone": row.phone,
        "manual": bool(row.manual),
        "opted_out": bool(row.opted_out),
        "quote": row.quote,
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }
